#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Pflegelöwen GmbH – Social-Media- & Content-Strategie (Kick-off Deck)."""

from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn

# ---------- Markenpalette ----------
PETROL   = RGBColor(0x12, 0x3B, 0x38)   # tiefes Pflege-Petrol (Anchor)
PETROL2  = RGBColor(0x1C, 0x55, 0x50)
GOLD     = RGBColor(0xC9, 0x97, 0x00)   # Löwen-Gold
GOLDLT   = RGBColor(0xE7, 0xC2, 0x5C)
CREAM    = RGBColor(0xFB, 0xF7, 0xEE)   # warmer Hintergrund
WHITE    = RGBColor(0xFF, 0xFF, 0xFF)
INK      = RGBColor(0x20, 0x2B, 0x2A)   # Fließtext
MUTE     = RGBColor(0x6A, 0x78, 0x76)   # Sekundärtext
CARD     = RGBColor(0xFF, 0xFF, 0xFF)
CARDLINE = RGBColor(0xE6, 0xDE, 0xCC)

FONT = "Calibri"
FONTH = "Calibri"

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
SW, SH = prs.slide_width, prs.slide_height
BLANK = prs.slide_layouts[6]

PAGENO = {"n": 0}

# ---------- Helfer ----------
def slide():
    return prs.slides.add_slide(BLANK)

def rect(s, x, y, w, h, fill, line=None, line_w=None, shape=MSO_SHAPE.RECTANGLE):
    sp = s.shapes.add_shape(shape, x, y, w, h)
    sp.shadow.inherit = False
    if fill is None:
        sp.fill.background()
    else:
        sp.fill.solid(); sp.fill.fore_color.rgb = fill
    if line is None:
        sp.line.fill.background()
    else:
        sp.line.color.rgb = line
        sp.line.width = line_w or Pt(1)
    return sp

def txt(s, x, y, w, h, runs, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP,
        space_after=4, line_spacing=1.0, wrap=True):
    """runs: list of paragraphs; each paragraph = list of (text, size, bold, color, font)."""
    tb = s.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame
    tf.word_wrap = wrap
    tf.vertical_anchor = anchor
    tf.margin_left = 0; tf.margin_right = 0; tf.margin_top = 0; tf.margin_bottom = 0
    for i, para in enumerate(runs):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.space_after = Pt(space_after)
        p.space_before = Pt(0)
        p.line_spacing = line_spacing
        for (t, sz, b, c, fn) in para:
            r = p.add_run(); r.text = t
            r.font.size = Pt(sz); r.font.bold = b
            r.font.color.rgb = c; r.font.name = fn
    return tb

def R(t, sz, b=False, c=INK, fn=FONT):
    return (t, sz, b, c, fn)

def bg(s, color=CREAM):
    rect(s, 0, 0, SW, SH, color)

def footer(s, label):
    PAGENO["n"] += 1
    rect(s, 0, SH - Inches(0.34), SW, Inches(0.34), PETROL)
    txt(s, Inches(0.55), SH - Inches(0.34), Inches(7), Inches(0.34),
        [[R("Pflegelöwen GmbH  ·  Social-Media- & Content-Strategie", 9, False, GOLDLT)]],
        anchor=MSO_ANCHOR.MIDDLE)
    txt(s, SW - Inches(4.55), SH - Inches(0.34), Inches(4), Inches(0.34),
        [[R(label + "   |   " + str(PAGENO["n"]).zfill(2), 9, False, CREAM)]],
        align=PP_ALIGN.RIGHT, anchor=MSO_ANCHOR.MIDDLE)

def header(s, kicker, title, label):
    bg(s)
    rect(s, 0, 0, Inches(0.22), SH, GOLD)               # Akzentkante links
    txt(s, Inches(0.55), Inches(0.42), Inches(11), Inches(0.3),
        [[R(kicker.upper(), 12, True, GOLD)]])
    txt(s, Inches(0.55), Inches(0.74), Inches(12.2), Inches(0.95),
        [[R(title, 30, True, PETROL)]])
    rect(s, Inches(0.58), Inches(1.62), Inches(1.1), Pt(3), GOLD)
    footer(s, label)

def chip(s, x, y, w, h, text, fill, tcolor, size=11):
    c = rect(s, x, y, w, h, fill, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    c.adjustments[0] = 0.5
    txt(s, x, y, w, h, [[R(text, size, True, tcolor)]],
        align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)

def card(s, x, y, w, h, head, lines, accent=GOLD, headsize=15, bodysize=11.5):
    rect(s, x, y, w, h, CARD, line=CARDLINE, line_w=Pt(1), shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    rect(s, x, y, Inches(0.10), h, accent, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    paras = [[R(head, headsize, True, PETROL)]]
    for ln in lines:
        paras.append([R("•  ", bodysize, True, accent), R(ln, bodysize, False, INK)])
    txt(s, x + Inches(0.28), y + Inches(0.18), w - Inches(0.45), h - Inches(0.3),
        paras, space_after=4, line_spacing=1.02)

# =========================================================
# 1 — TITEL
# =========================================================
s = slide()
rect(s, 0, 0, SW, SH, PETROL)
rect(s, 0, 0, SW, Inches(2.3), PETROL2)
# Löwen-Akzentband
rect(s, 0, SH - Inches(0.9), SW, Inches(0.9), GOLD)
rect(s, 0, SH - Inches(0.96), SW, Inches(0.06), GOLDLT)
# „Wortmarke" Löwe
txt(s, Inches(0.7), Inches(0.55), Inches(8), Inches(0.5),
    [[R("🦁  PFLEGELÖWEN", 18, True, GOLDLT)]])
txt(s, Inches(0.7), Inches(2.25), Inches(12), Inches(2.2),
    [[R("Social Media & Content", 46, True, WHITE)],
     [R("Strategie 2026", 46, True, GOLDLT)]], line_spacing=1.0)
txt(s, Inches(0.72), Inches(4.5), Inches(11.5), Inches(0.6),
    [[R("Vom Pflegedienst zur Lieblingsmarke der Region – nahbar, sichtbar, löwenstark.",
        16, False, RGBColor(0xD9,0xE5,0xE3))]])
txt(s, Inches(0.72), SH - Inches(0.83), Inches(9), Inches(0.7),
    [[R("Kick-off & Fahrplan  ·  Geschäftspartner: Hiko Krassevic", 12.5, True, PETROL)],
     [R("Stand: 15. Juni 2026", 11, False, PETROL)]], space_after=1)

# =========================================================
# 2 — AGENDA
# =========================================================
s = slide(); header(s, "Überblick", "Worüber wir heute sprechen", "Agenda")
items = [
    ("01", "Ausgangslage", "Wo Pflegelöwen heute steht"),
    ("02", "Zielgruppen", "Wen wir erreichen – und warum"),
    ("03", "Die Leitidee", "Unsere Strategie in einem Satz"),
    ("04", "Content-Säulen", "Worüber wir reden – 5 Themenwelten"),
    ("05", "Output & Format-Mix", "Was wir pro Woche liefern"),
    ("06", "Wochenplan Mo–So", "Der konkrete Sende-Rhythmus"),
    ("07", "Ideen-Backlog", "Konkrete Video- & Reel-Ideen"),
    ("08", "Trello-Workflow", "Wie wir produzieren & steuern"),
    ("09", "Ads & Neukunden", "Bezahlte Reichweite mit Budget"),
    ("10", "KPIs & Fahrplan", "Erfolg messen – 30/60/90 Tage"),
]
colw = Inches(6.05); rowh = Inches(0.84)
x0 = Inches(0.7); y0 = Inches(2.0); gx = Inches(0.25); gy = Inches(0.18)
for i, (n, t, d) in enumerate(items):
    col = i // 5; row = i % 5
    x = x0 + col * (colw + gx); y = y0 + row * (rowh + gy)
    rect(s, x, y, colw, rowh, CARD, line=CARDLINE, line_w=Pt(1), shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    rect(s, x, y, Inches(0.85), rowh, PETROL, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    txt(s, x, y, Inches(0.85), rowh, [[R(n, 18, True, GOLDLT)]],
        align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    txt(s, x + Inches(1.0), y, colw - Inches(1.1), rowh,
        [[R(t, 13.5, True, PETROL)], [R(d, 10.5, False, MUTE)]],
        anchor=MSO_ANCHOR.MIDDLE, space_after=1)

# =========================================================
# 3 — AUSGANGSLAGE
# =========================================================
s = slide(); header(s, "01 · Ausgangslage", "Wo Pflegelöwen heute steht", "Ausgangslage")
card(s, Inches(0.7), Inches(2.0), Inches(6.05), Inches(2.35), "Die Marke",
     ["Ambulanter Pflegedienst in Hannover – „Löwenstarke Pflege“",
      "Starkes Leistungsportfolio: Grund-/Behandlungspflege, 24/7, SAPV, Intensiv- & Tagespflege",
      "Klare Werte: Sicherheit, Würde, Lebensqualität – zu Hause bleiben",
      "Sympathisches Löwen-Branding mit Wiedererkennung"], accent=GOLD)
card(s, Inches(6.95), Inches(2.0), Inches(6.05), Inches(2.35), "Die Chance",
     ["Pflege wird im Netz „erklärt“ – aber selten emotional & nahbar erzählt",
      "Vorbild-Accounts (z. B. @stella_mom_ceo) zeigen: Gesicht + Geschichte wirkt",
      "Zwei Bedarfe gleichzeitig bespielbar: Neukund:innen UND Pflegekräfte",
      "Regional kaum Wettbewerb mit echtem Content-System"], accent=GOLD)
# Aufgabe-Banner
rect(s, Inches(0.7), Inches(4.6), Inches(12.3), Inches(1.95), PETROL, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
txt(s, Inches(1.05), Inches(4.8), Inches(11.7), Inches(0.4),
    [[R("UNSER AUFTRAG", 12, True, GOLDLT)]])
txt(s, Inches(1.05), Inches(5.18), Inches(11.7), Inches(1.3),
    [[R("„Nimm die Zügel in die Hand. Bring mich nach vorne – als wäre ich dein Kunde.“", 17, True, WHITE)],
     [R("Wir bauen ein wiederholbares Content-System, testen mit Budget und beweisen, was funktioniert. "
        "Wenn es funktioniert, ist es verkaufbar. Deshalb: erst die Craft, dann die Skalierung.",
        12.5, False, RGBColor(0xD9,0xE5,0xE3))]], space_after=6)

# =========================================================
# 4 — ZIELGRUPPEN
# =========================================================
s = slide(); header(s, "02 · Zielgruppen", "Drei Menschen, die wir erreichen", "Zielgruppen")
cards = [
    ("👵  Senior:innen & Angehörige", GOLD,
     ["Töchter/Söhne (45–65), die Hilfe für Eltern suchen",
      "Emotion: Überforderung, Schuldgefühl, Suche nach Vertrauen",
      "Wir liefern: Orientierung, Sicherheit, ein Gesicht zum Vertrauen"]),
    ("🧑‍⚕️  Pflegekräfte (Recruiting)", PETROL2,
     ["Fachkräfte & Quereinsteiger:innen in/um Hannover",
      "Emotion: Wertschätzung, faire Bedingungen, Teamgefühl",
      "Wir liefern: Einblicke ins Team, Kultur, „hier will ich arbeiten“"]),
    ("🤝  Multiplikatoren & Region", GOLDLT,
     ["Ärzt:innen, Kliniken, Sozialdienste, lokale Community",
      "Emotion: Verlässlichkeit, Reputation, Bekanntheit",
      "Wir liefern: Sichtbarkeit & Glaubwürdigkeit als Marke aus Hannover"]),
]
x = Inches(0.7); w = Inches(3.98); gap = Inches(0.18)
for i, (h, ac, lines) in enumerate(cards):
    xi = x + i * (w + gap)
    rect(s, xi, Inches(2.0), w, Inches(3.9), CARD, line=CARDLINE, line_w=Pt(1), shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    rect(s, xi, Inches(2.0), w, Inches(0.95), ac if ac != GOLDLT else GOLD, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    txt(s, xi + Inches(0.25), Inches(2.0), w - Inches(0.4), Inches(0.95),
        [[R(h, 14.5, True, WHITE)]], anchor=MSO_ANCHOR.MIDDLE)
    paras = []
    for ln in lines:
        paras.append([R("›  ", 12, True, GOLD), R(ln, 11.5, False, INK)])
    txt(s, xi + Inches(0.28), Inches(3.15), w - Inches(0.5), Inches(2.6), paras, space_after=8, line_spacing=1.03)
txt(s, Inches(0.7), Inches(6.1), Inches(12.3), Inches(0.6),
    [[R("Fokus Phase 1: ", 12.5, True, PETROL),
      R("Neukund:innen-Vertrauen + Recruiting parallel – beides zahlt auf dieselbe Marke ein.", 12.5, False, INK)]])

# =========================================================
# 5 — LEITIDEE
# =========================================================
s = slide()
rect(s, 0, 0, SW, SH, PETROL)
rect(s, 0, 0, Inches(0.22), SH, GOLD)
txt(s, Inches(0.7), Inches(0.7), Inches(11), Inches(0.4), [[R("03 · DIE LEITIDEE", 13, True, GOLDLT)]])
txt(s, Inches(0.7), Inches(1.7), Inches(12), Inches(2.6),
    [[R("Wir machen Pflege", 30, True, WHITE)],
     [R("menschlich sichtbar.", 40, True, GOLDLT)]], line_spacing=1.05)
txt(s, Inches(0.72), Inches(3.9), Inches(11.8), Inches(1.0),
    [[R("Statt Logos und Leistungslisten zeigen wir Gesichter, Geschichten und Momente – "
       "den Alltag des Pflegelöwen-Teams. So entsteht Vertrauen, bevor jemand anruft.", 15, False, RGBColor(0xDB,0xE6,0xE4))]])
pills = ["Nahbar statt steril", "Gesicht statt Logo", "Geschichte statt Werbung", "Konstanz statt Zufall"]
px = Inches(0.7)
for p in pills:
    w = Inches(0.16 + 0.115 * len(p))
    chip(s, px, Inches(5.25), w, Inches(0.52), p, GOLD, PETROL, size=12)
    px += w + Inches(0.22)
txt(s, Inches(0.7), Inches(6.25), Inches(12), Inches(0.6),
    [[R("Benchmark-Tonalität: @stella_mom_ceo – persönlich, ehrlich, mit Haltung. Genau dieses Gefühl bauen wir für Pflegelöwen.",
        11.5, False, RGBColor(0xB9,0xCC,0xC9))]])

# =========================================================
# 6 — CONTENT-SÄULEN
# =========================================================
s = slide(); header(s, "04 · Content-Säulen", "Fünf Themenwelten, die wir rotieren", "Content-Säulen")
pillars = [
    ("Vertrauen & Kompetenz", "Pflegewissen einfach erklärt: Pflegegrad, Anträge, Tipps für Angehörige.", "Spart Anruf-Hürden"),
    ("Menschen & Team", "Mitarbeiter:innen vorstellen, Behind-the-Scenes, „Ein Tag mit…“.", "Recruiting + Nähe"),
    ("Herz & Momente", "Berührende Geschichten von Senior:innen, kleine Alltagssiege.", "Emotion & Shares"),
    ("Service & Beweis", "Leistungen zeigen, Bewertungen, Vorher/Nachher der Betreuung.", "Conversion"),
    ("Region Hannover", "Lokale Bezüge, Aktionen, „Pflegelöwen vor Ort“.", "Lokale Reichweite"),
]
x0 = Inches(0.7); y0 = Inches(2.0); w = Inches(2.42); h = Inches(3.95); gap = Inches(0.12)
for i, (t, d, tag) in enumerate(pillars):
    xi = x0 + i * (w + gap)
    rect(s, xi, y0, w, h, CARD, line=CARDLINE, line_w=Pt(1), shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    rect(s, xi, y0, w, Inches(0.55), GOLD, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    txt(s, xi, y0 + Inches(0.02), w, Inches(0.55), [[R("0"+str(i+1), 16, True, PETROL)]],
        align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    txt(s, xi + Inches(0.2), y0 + Inches(0.75), w - Inches(0.36), Inches(1.1),
        [[R(t, 13, True, PETROL)]], anchor=MSO_ANCHOR.TOP, line_spacing=1.0)
    txt(s, xi + Inches(0.2), y0 + Inches(1.65), w - Inches(0.36), Inches(1.7),
        [[R(d, 10.5, False, INK)]], line_spacing=1.04)
    rect(s, xi + Inches(0.2), y0 + h - Inches(0.62), w - Inches(0.4), Inches(0.42),
         CREAM, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    txt(s, xi + Inches(0.2), y0 + h - Inches(0.62), w - Inches(0.4), Inches(0.42),
        [[R(tag, 9.5, True, GOLD)]], align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
txt(s, Inches(0.7), Inches(6.15), Inches(12.3), Inches(0.5),
    [[R("Rotation sorgt für Abwechslung – jede Säule erscheint regelmäßig, kein Thema dominiert.", 11.5, False, MUTE)]])

# =========================================================
# 7 — OUTPUT / FORMAT-MIX
# =========================================================
s = slide(); header(s, "05 · Output", "Was wir pro Woche liefern", "Output")
stats = [
    ("2", "Videos / Woche", "Storytelling & Aufklärung (30–90 Sek.)"),
    ("2", "Reels / Woche", "Trend- & Hook-getrieben für Reichweite"),
    ("3–4", "Beiträge / Stories", "Karussells, Tipps, Team, Umfragen"),
    ("7/7", "Tage aktiv", "Community-Management & Antworten"),
]
x0 = Inches(0.7); w = Inches(2.98); gap = Inches(0.17)
for i, (big, lab, sub) in enumerate(stats):
    xi = x0 + i * (w + gap)
    rect(s, xi, Inches(2.05), w, Inches(2.15), PETROL, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    txt(s, xi, Inches(2.2), w, Inches(0.95), [[R(big, 44, True, GOLDLT)]],
        align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    txt(s, xi, Inches(3.15), w, Inches(0.4), [[R(lab, 13, True, WHITE)]], align=PP_ALIGN.CENTER)
    txt(s, xi + Inches(0.2), Inches(3.55), w - Inches(0.4), Inches(0.6),
        [[R(sub, 10, False, RGBColor(0xC7,0xD6,0xD4))]], align=PP_ALIGN.CENTER, line_spacing=1.0)
card(s, Inches(0.7), Inches(4.45), Inches(6.05), Inches(2.05), "Qualität vor Menge",
     ["Lieber 2 starke Videos als 6 belanglose – „work on the craft“",
      "Wiederverwertung: 1 Dreh → mehrere Schnitte (Reel, Story, Karussell)",
      "Jeder Post hat ein Ziel: Reichweite, Vertrauen oder Conversion"], accent=GOLD)
card(s, Inches(6.95), Inches(4.45), Inches(6.05), Inches(2.05), "Effizienter Dreh-Rhythmus",
     ["1 Drehtag alle 2 Wochen = Vorrat für ~2 Wochen Content",
      "Smartphone-first + gutes Licht/Ton statt teurer Produktion",
      "Hiko als Gesicht/Stimme einbinden, wo es passt"], accent=PETROL2)

# =========================================================
# 8 — WOCHENPLAN
# =========================================================
s = slide(); header(s, "06 · Wochenplan", "Unser Sende-Rhythmus Mo – So", "Wochenplan")
days = [
    ("MO", "Reel", "Hook/Trend", "Reichweite", GOLD),
    ("DI", "Karussell", "Pflege-Tipp", "Vertrauen", PETROL2),
    ("MI", "Video", "Team/Story", "Recruiting", GOLD),
    ("DO", "Story-Reihe", "Q&A / Umfrage", "Community", PETROL2),
    ("FR", "Reel", "Emotion/Moment", "Shares", GOLD),
    ("SA", "Video", "Service/Beweis", "Conversion", PETROL2),
    ("SO", "Beitrag", "Rückblick/Zitat", "Bindung", GOLD),
]
x0 = Inches(0.62); y0 = Inches(2.05); w = Inches(1.74); gap = Inches(0.05); h = Inches(3.55)
for i, (d, fmt, topic, goal, ac) in enumerate(days):
    xi = x0 + i * (w + gap)
    rect(s, xi, y0, w, h, CARD, line=CARDLINE, line_w=Pt(1), shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    rect(s, xi, y0, w, Inches(0.6), ac, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    txt(s, xi, y0, w, Inches(0.6), [[R(d, 16, True, WHITE if ac==PETROL2 else PETROL)]],
        align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    txt(s, xi + Inches(0.12), y0 + Inches(0.78), w - Inches(0.24), Inches(0.7),
        [[R(fmt, 12.5, True, PETROL)]], align=PP_ALIGN.CENTER, line_spacing=0.95)
    txt(s, xi + Inches(0.1), y0 + Inches(1.55), w - Inches(0.2), Inches(0.9),
        [[R(topic, 10.5, False, INK)]], align=PP_ALIGN.CENTER, line_spacing=1.0)
    rect(s, xi + Inches(0.16), y0 + h - Inches(0.55), w - Inches(0.32), Inches(0.4),
         CREAM, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    txt(s, xi + Inches(0.1), y0 + h - Inches(0.55), w - Inches(0.2), Inches(0.4),
        [[R(goal, 9, True, GOLD)]], align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
txt(s, Inches(0.62), Inches(5.8), Inches(12.4), Inches(0.7),
    [[R("Fixe Slots = Verlässlichkeit für den Algorithmus und fürs Team.  ", 11.5, True, PETROL),
      R("Beste Zeiten: werktags 7–8 & 18–20 Uhr, So vormittags – wird datenbasiert nachjustiert.", 11.5, False, INK)]])

# =========================================================
# 9 — IDEEN-BACKLOG
# =========================================================
s = slide(); header(s, "07 · Ideen-Backlog", "Konkrete Video- & Reel-Ideen für den Start", "Ideen")
ideas_left = [
    "„3 Dinge, die jede:r über Pflegegrade wissen muss“",
    "Ein Tag mit unserer Pflegekraft Maria (Mini-Doku)",
    "„Was Angehörige am meisten falsch machen“ (Hook)",
    "Senior:in erzählt: schönster Moment dieser Woche",
    "Mythos vs. Fakt: „Pflege zu Hause ist unbezahlbar?“",
]
ideas_right = [
    "Team stellt sich vor: „Warum ich Pflegelöwe wurde“",
    "Wohnungs-Check: 5 Stolperfallen für Senior:innen",
    "Behind-the-Scenes: So startet unser Team in den Tag",
    "Bewertung vorgelesen + Reaktion (Social Proof)",
    "„Wir suchen dich!“ – Recruiting-Reel mit Haltung",
]
def idea_col(x, items, start):
    rect(s, x, Inches(2.0), Inches(6.05), Inches(4.35), CARD, line=CARDLINE, line_w=Pt(1), shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    paras = []
    for j, it in enumerate(items):
        paras.append([R(f"{start+j:02d}  ", 12.5, True, GOLD), R(it, 12.5, False, INK)])
        paras.append([R("", 4, False, INK)])
    txt(s, x + Inches(0.35), Inches(2.28), Inches(5.5), Inches(3.9), paras, space_after=6, line_spacing=1.05)
idea_col(Inches(0.7), ideas_left, 1)
idea_col(Inches(6.95), ideas_right, 6)
txt(s, Inches(0.7), Inches(6.5), Inches(12.3), Inches(0.4),
    [[R("Alle Ideen leben im Trello-Board „Ideen“ und werden wöchentlich nachgefüllt und priorisiert.", 11, False, MUTE)]])

# =========================================================
# 10 — TRELLO-WORKFLOW
# =========================================================
s = slide(); header(s, "08 · Workflow", "Unser Trello-Board: von der Idee zum Post", "Workflow")
cols = [
    ("💡 IDEEN", "Sammelbecken für alle Einfälle & Trends", GOLD),
    ("🎬 IN ARBEIT", "Skript, Dreh, Schnitt – wer macht was", PETROL2),
    ("👀 REVIEW", "Freigabe durch Hiko / Qualitätscheck", GOLD),
    ("📅 GEPLANT", "Caption + Hashtags + Termin gesetzt", PETROL2),
    ("✅ LIVE", "Veröffentlicht – Performance wird getrackt", GOLD),
]
x0 = Inches(0.62); y0 = Inches(2.15); w = Inches(2.32); gap = Inches(0.18); h = Inches(2.9)
for i, (t, d, ac) in enumerate(cols):
    xi = x0 + i * (w + gap)
    rect(s, xi, y0, w, h, CARD, line=CARDLINE, line_w=Pt(1), shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    rect(s, xi, y0, w, Inches(0.62), ac, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    txt(s, xi, y0, w, Inches(0.62), [[R(t, 12, True, PETROL if ac==GOLD else WHITE)]],
        align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    txt(s, xi + Inches(0.2), y0 + Inches(0.85), w - Inches(0.4), Inches(1.9),
        [[R(d, 11, False, INK)]], align=PP_ALIGN.CENTER, line_spacing=1.08)
    if i < 4:
        ar = rect(s, xi + w - Inches(0.02), y0 + h/2 - Inches(0.18), gap + Inches(0.04), Inches(0.36),
                  GOLD, shape=MSO_SHAPE.CHEVRON)
rect(s, Inches(0.62), Inches(5.45), Inches(12.4), Inches(1.05), PETROL, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
txt(s, Inches(0.95), Inches(5.58), Inches(11.8), Inches(0.85),
    [[R("Transparenz für Hiko: ", 12.5, True, GOLDLT),
      R("Jederzeit sichtbar, was in Arbeit ist und was als Nächstes kommt. "
        "Wöchentliches 15-Min-Sync zur Freigabe & Feedback.", 12.5, False, WHITE)]], line_spacing=1.05)

# =========================================================
# 11 — ADS / NEUKUNDEN
# =========================================================
s = slide(); header(s, "09 · Ads & Neukunden", "Bezahlte Reichweite – kontrolliert testen", "Ads")
card(s, Inches(0.7), Inches(2.0), Inches(6.05), Inches(2.2), "Ziel der Ads",
     ["Neukunden-Anfragen (Beratungstermine) generieren",
      "Parallel: Bewerbungen von Pflegekräften",
      "Beste organische Inhalte als Ad skalieren („was wirkt, boosten“)"], accent=GOLD)
card(s, Inches(6.95), Inches(2.0), Inches(6.05), Inches(2.2), "Setup",
     ["Meta (Instagram/Facebook) als Start – Region Hannover + Umland",
      "Lead-Formular & Klick-zu-WhatsApp/Anruf",
      "Saubere Conversion-Messung (Pixel/Events) von Tag 1"], accent=PETROL2)
# Budget-Stufen
rect(s, Inches(0.7), Inches(4.45), Inches(12.3), Inches(2.05), CREAM, line=CARDLINE, line_w=Pt(1), shape=MSO_SHAPE.ROUNDED_RECTANGLE)
txt(s, Inches(0.95), Inches(4.6), Inches(11), Inches(0.4), [[R("TEST-BUDGET IN STUFEN", 12, True, GOLD)]])
steps = [
    ("Phase 1 · Test", "klein & breit", "2–3 Creatives, Zielgruppen & Botschaften testen"),
    ("Phase 2 · Lernen", "Gewinner finden", "Bestes Creative + beste Zielgruppe identifizieren"),
    ("Phase 3 · Skalieren", "Budget aufs Beste", "Was funktioniert, bekommt mehr Budget"),
]
x = Inches(0.95)
for i,(t,b,d) in enumerate(steps):
    w = Inches(3.85)
    xi = Inches(0.95) + i*(w+Inches(0.18))
    rect(s, xi, Inches(5.05), w, Inches(1.25), WHITE, line=CARDLINE, line_w=Pt(1), shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    txt(s, xi+Inches(0.22), Inches(5.18), w-Inches(0.4), Inches(1.0),
        [[R(t, 12.5, True, PETROL)],[R(b, 11, True, GOLD)],[R(d, 10, False, INK)]], space_after=2, line_spacing=1.0)

# =========================================================
# 12 — KPIs
# =========================================================
s = slide(); header(s, "10 · Erfolg messen", "Woran wir erkennen, dass es funktioniert", "KPIs")
kpis = [
    ("Reichweite & Views", "Werden wir gesehen?", "Aufrufe, Reichweite, Follower-Wachstum"),
    ("Engagement", "Berührt es Menschen?", "Saves, Shares, Kommentare, Watchtime"),
    ("Profil → Aktion", "Werden sie neugierig?", "Profilbesuche, Link-Klicks, Nachrichten"),
    ("Anfragen / Leads", "Kommen Kund:innen?", "Beratungsanfragen, Anrufe, Bewerbungen"),
]
x0 = Inches(0.7); w = Inches(2.98); gap = Inches(0.17)
for i,(t,q,d) in enumerate(kpis):
    xi = x0 + i*(w+gap)
    rect(s, xi, Inches(2.0), w, Inches(2.55), CARD, line=CARDLINE, line_w=Pt(1), shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    rect(s, xi, Inches(2.0), w, Inches(0.12), GOLD, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    txt(s, xi+Inches(0.22), Inches(2.28), w-Inches(0.4), Inches(0.7), [[R(t, 13, True, PETROL)]], line_spacing=0.95)
    txt(s, xi+Inches(0.22), Inches(3.05), w-Inches(0.4), Inches(0.5), [[R(q, 11, True, GOLD)]])
    txt(s, xi+Inches(0.22), Inches(3.5), w-Inches(0.4), Inches(0.9), [[R(d, 10.5, False, INK)]], line_spacing=1.05)
rect(s, Inches(0.7), Inches(4.8), Inches(12.3), Inches(1.7), PETROL, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
txt(s, Inches(1.0), Inches(4.98), Inches(11.6), Inches(0.4), [[R("DATENGETRIEBEN ARBEITEN", 12, True, GOLDLT)]])
txt(s, Inches(1.0), Inches(5.36), Inches(11.6), Inches(1.0),
    [[R("Monatliches Reporting in Klartext: Was lief gut, was nicht, was testen wir als Nächstes. "
       "Wir entscheiden nach Zahlen – nicht nach Bauchgefühl. So sehen wir gemeinsam: „Kommen wir an oder nicht?“",
       13, False, WHITE)]], line_spacing=1.08)

# =========================================================
# 13 — FAHRPLAN 30/60/90
# =========================================================
s = slide(); header(s, "Fahrplan", "Die nächsten 90 Tage", "Fahrplan")
phases = [
    ("TAG 0–30", "Fundament", GOLD,
     ["Profil schärfen (Bio, Highlights, Look)",
      "Erster Drehtag + Content-Vorrat",
      "Trello-Board scharf schalten",
      "Sende-Rhythmus Mo–So starten"]),
    ("TAG 31–60", "Lernen & Testen", PETROL2,
     ["Erste Ads-Tests live schalten",
      "Datenanalyse: was funktioniert?",
      "Formate doppeln, die ziehen",
      "Erste Leads/Bewerbungen tracken"]),
    ("TAG 61–90", "Skalieren", GOLD,
     ["Budget auf Gewinner-Creatives",
      "Wiederholbares System dokumentieren",
      "Reporting & klare KPIs",
      "Bereit, das Modell zu verkaufen"]),
]
x0 = Inches(0.7); w = Inches(3.98); gap = Inches(0.18)
for i,(tag,t,ac,lines) in enumerate(phases):
    xi = x0 + i*(w+gap)
    rect(s, xi, Inches(2.0), w, Inches(4.3), CARD, line=CARDLINE, line_w=Pt(1), shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    rect(s, xi, Inches(2.0), w, Inches(1.0), ac if ac==PETROL2 else PETROL, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    txt(s, xi+Inches(0.25), Inches(2.12), w-Inches(0.4), Inches(0.45), [[R(tag, 12, True, GOLDLT)]])
    txt(s, xi+Inches(0.25), Inches(2.5), w-Inches(0.4), Inches(0.5), [[R(t, 17, True, WHITE)]])
    paras = [[R("•  ", 12, True, ac), R(ln, 12, False, INK)] for ln in lines]
    txt(s, xi+Inches(0.3), Inches(3.25), w-Inches(0.55), Inches(2.9), paras, space_after=9, line_spacing=1.05)

# =========================================================
# 14 — CLOSING
# =========================================================
s = slide()
rect(s, 0, 0, SW, SH, PETROL)
rect(s, 0, SH-Inches(0.9), SW, Inches(0.9), GOLD)
rect(s, 0, SH-Inches(0.96), SW, Inches(0.06), GOLDLT)
txt(s, Inches(0.7), Inches(0.7), Inches(8), Inches(0.5), [[R("🦁  PFLEGELÖWEN", 16, True, GOLDLT)]])
txt(s, Inches(0.7), Inches(2.2), Inches(12), Inches(2.2),
    [[R("Erst die Craft.", 42, True, WHITE)],
     [R("Dann die Skalierung.", 42, True, GOLDLT)]], line_spacing=1.02)
txt(s, Inches(0.72), Inches(4.45), Inches(11.5), Inches(1.0),
    [[R("Wir bauen ein System, das funktioniert – beweisen es mit Daten – und machen es verkaufbar. "
       "Mit Verlässlichkeit, Herz und Löwenmut. Und ja: Lass uns Spaß dabei haben.",
       15, False, RGBColor(0xDB,0xE6,0xE4))]])
txt(s, Inches(0.72), SH-Inches(0.82), Inches(11), Inches(0.7),
    [[R("Nächster Schritt: Board freigeben, ersten Drehtag terminieren, Ads-Test vorbereiten.", 12.5, True, PETROL)]])

prs.save("/home/user/Pflegeloewen_Social_Media_Strategie.pptx")
print("OK – Folien:", len(prs.slides._sldIdLst))
